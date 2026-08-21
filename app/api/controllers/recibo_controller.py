from app.api.services.recibo_service import get_propiedades_ajustar, actualizar_fechas, actualizar_ipc,get_ipc
from app.api.services.libro_diario_service import marcar_todas_adeuda
from openpyxl import load_workbook

def get_recibos_ajustar(mes_liquidacion: int, año_liquidacion: int) -> list:
    try:
        return get_propiedades_ajustar(mes_liquidacion, año_liquidacion)
    except Exception as e:
        raise e

def actualizar_recibos(mes_liquidacion: int, año_liquidacion: int):
    recibos_ajustar = get_propiedades_ajustar(mes_liquidacion, año_liquidacion) #obtengo que recibos debo ajustar
    
    wb = load_workbook("templates/RECIBO INMOBILIARIO.xlsx") #Carga el excel de recibos

    if  recibos_ajustar:
        valores_ipc=get_ipc() #Obtengo los ultimos 4 valores del ipc
    
    #recorre todas las hojas del excel
    for recibo in wb.sheetnames:
        actualizar_fechas(recibo, mes_liquidacion, año_liquidacion, wb)  # Actualiza las fechas en el recibo

        if recibo in recibos_ajustar: #Si tiene que actualizar
            actualizar_ipc(recibo, wb, valores_ipc)  # Actualiza el IPC en el recibo

    wb.save("templates/RECIBO INMOBILIARIO.xlsx")

    # Empieza un período nuevo: lo cobrado el mes pasado ya no cuenta.
    marcar_todas_adeuda()
