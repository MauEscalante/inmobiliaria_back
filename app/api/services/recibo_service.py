from datetime import date
from decimal import Decimal

import requests
from openpyxl import load_workbook
from openpyxl.workbook.workbook import Workbook
from sqlalchemy import text

from app.config import settings
from app.database.connection import SessionLocal
from app.utils.helpers import serializar_fila

MESES = {
    1: "Enero",
    2: "Febrero",
    3: "Marzo",
    4: "Abril",
    5: "Mayo",
    6: "Junio",
    7: "Julio",
    8: "Agosto",
    9: "Septiembre",
    10: "Octubre",
    11: "Noviembre",
    12: "Diciembre",
}

# Solo las columnas que el recurso expone, en vez del SELECT * que devolvía la
# tabla entera de contratos sin contrato de datos.
#
# El CASE traduce la periodicidad a meses acá porque la decisión se toma en SQL;
# la versión Python de ese mapa es contrato.MESES_POR_PERIODICIDAD, que usa el
# cálculo de rescisión. Si se agrega una periodicidad hay que tocar las dos.
CONTRATOS_A_AJUSTAR_SELECT = """
    SELECT contrato_id, propiedad, fecha_inicio, fecha_fin,
           importe_inicial, periodicidad, tipo_ajuste
    FROM contrato
    WHERE tipo_ajuste = 'IPC'
      AND estado = 'Activo'
      AND TIMESTAMPDIFF(MONTH, fecha_inicio, :fecha_liquidacion) > 0
      AND MOD(
              TIMESTAMPDIFF(MONTH, fecha_inicio, :fecha_liquidacion),
              CASE periodicidad
                  WHEN 'Trimestral' THEN 3
                  WHEN 'Cuatrimestral' THEN 4
                  WHEN 'Semestral' THEN 6
              END
          ) = 0
    ORDER BY contrato_id
"""


def get_propiedades_ajustar(mes_liquidacion: int, anio_liquidacion: int) -> list:
    """Contratos a los que les toca ajuste por IPC en el período pedido."""
    db = SessionLocal()
    try:
        fecha_liquidacion = date(anio_liquidacion, mes_liquidacion, 1)
        filas = db.execute(
            text(CONTRATOS_A_AJUSTAR_SELECT),
            {"fecha_liquidacion": fecha_liquidacion},
        ).mappings().all()
        return [serializar_fila(fila) for fila in filas]
    finally:
        db.close()


def cargar_planilla() -> Workbook:
    """Planilla completa, en modo lectura/escritura. Es cara: ~40s con 120 hojas.

    Usar solo cuando haya que modificar y guardar; para leer metadatos está
    nombres_de_hojas().
    """
    return load_workbook(settings.RECIBOS_TEMPLATE)


def nombres_de_hojas() -> list[str]:
    """Nombres de las hojas de la planilla.

    Se abre en read_only porque el workbook completo parsea todas las celdas y
    tarda ~40 segundos; en modo lectura son décimas de segundo.
    """
    wb = load_workbook(settings.RECIBOS_TEMPLATE, read_only=True)
    try:
        return list(wb.sheetnames)
    finally:
        wb.close()


def actualizar_fechas(
    recibo: str, mes_liquidacion: int, anio_liquidacion: int, wb: Workbook
) -> None:
    ws = wb[recibo]
    # cambia la fecha del recibo y el mes
    ws["I13"].value = mes_liquidacion
    ws["C22"].value = MESES[mes_liquidacion]
    ws["J13"].value = anio_liquidacion


def factor_ipc(valores_ipc: list) -> Decimal:
    """Cuánto se multiplica el alquiler con los valores de IPC del período."""
    factor = Decimal("1")
    for valor in valores_ipc:
        factor *= Decimal(str(valor["valor"]))
    return factor


def actualizar_importe(recibo: str, wb: Workbook, importe) -> None:
    """Escribe en el recibo el alquiler que quedó guardado en valor_historico.

    Antes se multiplicaba la celda por sí misma, y como el recibo ya trae una
    aproximación hecha el valor se iba separando del de la base ajuste tras ajuste.
    La base manda: es de donde sale el importe del próximo re-ajuste y el de la
    penalidad por rescisión.
    """
    ws = wb[recibo]
    ws["E22"].value = float(importe)


def get_ipc() -> list:
    """Últimos 4 valores del IPC.

    La llamada lleva timeout y verificación de status: sin eso, una caída del
    servicio externo dejaba el request colgado.
    """
    response = requests.get(settings.IPC_API_URL, timeout=settings.IPC_TIMEOUT)
    response.raise_for_status()
    data = response.json().get("data", [])
    return data[-4:]
